"""Read the dividend workbook and reports without modifying source files."""

import argparse
import datetime as dt
import json
from pathlib import Path
import re

from .calculations import number


def timestamp():
    return dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(timespec="seconds")


def date_string(value):
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    try:
        return dt.date.fromisoformat(str(value)[:10]).isoformat()
    except ValueError:
        return None


def load_optional(path, fallback):
    if not path.is_file():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


def extract(source):
    import openpyxl

    source = Path(source).resolve()
    workbook_path = source / "data" / "daily-dy.xlsx"
    cache = load_optional(source / "scripts" / ".dy_cache.json", {})
    reports = sorted((source / "pictures").glob("[0-9]*/dividend_percentiles.json"))
    records = load_optional(reports[-1], []) if reports else []
    by_code = {str(row["code"]).zfill(6): row for row in records}
    quote_date = date_string(cache.get("ashare_prices", {}).get("date"))
    imported_at = timestamp()
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    stocks = []
    sector = "其他"
    try:
        for row in workbook["股票数据"].iter_rows(min_row=2, values_only=True):
            if not row[1] or row[2] is None:
                continue
            code = str(row[2]).strip().zfill(6)
            if not re.fullmatch(r"\d{6}", code):
                continue
            sector = str(row[0]).strip() if row[0] else sector
            report = by_code.get(code, {})
            annual = report.get("annual") or {}
            history = report.get("percentile_history") or {}
            unverified = str(row[13] if len(row) > 13 else "") == "待核实" or bool(report.get("error"))
            annual_value = number(row[4])
            annual_year = int(row[13]) if len(row) > 13 and number(row[13]) is not None else annual.get("year")
            rolling_value = number(report.get("total"))
            notes = []
            if annual and annual_value is not None and number(annual.get("total")) != annual_value:
                history = {}
                annual = {}
                rolling_value = None
                notes.append("工作簿年度股息与历史报告不一致，历史对比及滚动金额暂不可用；请运行原项目更新后重导入")
            if unverified:
                annual_value = None
                annual_year = None
                notes.append(report.get("error") or "年度分红待核实，暂不计算年度股息率")
            if history.get("error"):
                notes.append(history["error"])
            if history.get("warnings"):
                notes.extend(history["warnings"])
            if history.get("limitation") and history.get("samples"):
                notes.append(history["limitation"])
            stocks.append({
                "market": "A", "code": code, "name": str(row[1]).strip(), "sector": sector,
                "price": number(row[3]), "annual_dividend": annual_value,
                "rolling_dividend": rolling_value, "target_yield": number(row[7]),
                "confirmed": bool(annual.get("confirmed", row[5])) and not unverified,
                "annual_year": annual_year, "year_start_price": number(row[9]),
                "change_pct": number(cache.get("ashare_pct_chg", {}).get("data", {}).get(code)),
                "ex_dividend_date": date_string(row[11]), "quote_date": quote_date if number(cache.get("ashare_prices", {}).get("data", {}).get(code)) == number(row[3]) else None,
                "updated_at": imported_at, "a_code": None, "connect_factor": 0.8,
                "percentile_start": history.get("start"), "percentile_end": history.get("end"),
                "dividends": annual.get("parts") or [], "history": history.get("samples") or [],
                "source": report.get("source") or "本地 daily-dy.xlsx；行情：腾讯，分红：现有工作簿",
                "verification_note": "；".join(notes) or None,
            })
        fx_rate = None
        if "港股数据" in workbook.sheetnames:
            sheet = workbook["港股数据"]
            first = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            if "汇率" in first:
                label_index = first.index("汇率")
                fx_rate = number(first[label_index + 1]) if label_index + 1 < len(first) else None
            a_stocks = {stock["code"]: stock for stock in stocks}
            hk_ts = cache.get("hk_prices", {}).get("ts")
            hk_date = dt.datetime.fromtimestamp(hk_ts, dt.timezone(dt.timedelta(hours=8))).date().isoformat() if isinstance(hk_ts, (int, float)) else None
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0] or row[1] is None:
                    continue
                code = str(row[1]).strip().zfill(5)
                if not re.fullmatch(r"\d{5}", code):
                    continue
                a_code = str(row[2]).strip().zfill(6) if row[2] is not None else None
                a_stock = a_stocks.get(a_code, {})
                hk_annual = a_stock.get("annual_dividend") if a_stock else number(row[5])
                stocks.append({
                    "market": "HK", "code": code, "name": str(row[0]).strip(),
                    "sector": a_stock.get("sector", "其他"), "price": number(row[3]),
                    "annual_dividend": hk_annual, "rolling_dividend": a_stock.get("rolling_dividend"),
                    "target_yield": a_stock.get("target_yield"), "confirmed": a_stock.get("confirmed", False),
                    "annual_year": a_stock.get("annual_year"), "year_start_price": None,
                    "change_pct": None, "ex_dividend_date": date_string(row[12]),
                    "quote_date": None, "updated_at": imported_at, "a_code": a_code,
                    "connect_factor": number(row[11]),
                    "a_price": number(row[4]), "percentile_start": None, "percentile_end": None,
                    "dividends": a_stock.get("dividends", []), "history": [],
                    "source": "本地港股数据表；H股价格 HKD，年度每股股息 RMB",
                    "verification_note": "税后口径沿用工作簿系数；港股分红日期来源仅含已发生记录，不含未来排期。"
                                         + (f"港股缓存采集于 {hk_date}，未记录真实交易日期；刷新行情后可补全。" if hk_date else "港股原始缓存未记录真实交易日期；刷新行情后可补全。")
                                         + ("对应 A 股年度股息待核实，暂不计算港股年度股息率。" if a_stock and hk_annual is None else ""),
                })
    finally:
        workbook.close()
    if not stocks:
        raise ValueError("工作簿中没有可导入的有效股票")
    return {"stocks": stocks, "meta": {"quote_date": quote_date, "imported_at": imported_at, "last_refresh": None,
            "fx_rate": fx_rate, "source": "本地 daily-dy.xlsx、dividend_percentiles.json 与腾讯行情缓存"}}


def main():
    parser = argparse.ArgumentParser(description="从现有分红项目导出独立网站种子数据")
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", default=str(Path(__file__).parent / "data" / "seed.json"))
    args = parser.parse_args()
    seed = extract(args.source)
    destination = Path(args.output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(seed, ensure_ascii=False, separators=(",", ":"), allow_nan=False), encoding="utf-8")
    print(f"已导出 {len(seed['stocks'])} 支股票，行情日期 {seed['meta']['quote_date']} → {destination}")


if __name__ == "__main__":
    main()
