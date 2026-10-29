import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from backend.importer import extract


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        (self.root / "data").mkdir()
        (self.root / "scripts").mkdir()
        (self.root / "pictures" / "20261020").mkdir(parents=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "股票数据"
        sheet.append(["板块", "股票名", "股票代码", "当前价格", "最新年度每股股息", "分红实施确认", "最新股息率",
                      "目标股息率", "目标价格", "26初价格", "26以来涨跌", "最近除息日", "最新滚动年度股息率", "股息所属年度"])
        sheet.append(["银行", "有效标的", "600001", 10, 0.5, 1, 0.05, 0.05, 10, 8, 0.25, "2026-11-01", 0.06, 2025])
        sheet.append([None, "待核实标的", "600002", 8, 0.3, 0, None, None, None, None, None, None, None, "待核实"])
        hks = workbook.create_sheet("港股数据")
        hks.append(["股票名", "H股代码", "A股代码", "H股价格HKD", "A股价格CNY", "最新年度分红（RMB）",
                    "A/H溢价率", "A股股息率", "港股股息率", "港户股息率", "港股通股息率", "港股通税率", "最近除息日", "汇率", 0.8])
        hks.append(["有效港股", "00001", "600001", 8, 10, 0.5, None, None, None, None, None, 0.72, "2026-07-01"])
        workbook.save(self.root / "data" / "daily-dy.xlsx")
        self.report = [{"code": "600001", "name": "有效标的", "total": 0.6, "annual": {"year": 2025, "total": 0.5,
                        "confirmed": 1, "parts": [{"period": "2025-12-31", "dps": 0.5, "status": "实施分配"}]},
                        "percentile_history": {"start": "2021-10-20", "end": "2026-10-20", "samples": [{"date": "2026-10-20", "yield": 0.05, "close": 10}]}}]
        self.report_path = self.root / "pictures" / "20261020" / "dividend_percentiles.json"
        self.report_path.write_text(json.dumps(self.report), encoding="utf-8")
        (self.root / "scripts" / ".dy_cache.json").write_text(json.dumps({"ashare_prices": {"date": "2026-10-20"},
               "ashare_pct_chg": {"data": {"600001": -0.012}}, "hk_prices": {"ts": 1792483200}}), encoding="utf-8")

    def tearDown(self):
        self.directory.cleanup()

    def test_read_only_import_uses_real_values_and_inherits_sectors(self):
        before = (self.root / "data" / "daily-dy.xlsx").read_bytes()
        result = extract(self.root)
        self.assertEqual(len(result["stocks"]), 3)
        self.assertEqual([stock["sector"] for stock in result["stocks"]], ["银行", "银行", "银行"])
        self.assertIsNone(result["stocks"][1]["annual_dividend"])
        self.assertEqual(result["stocks"][0]["rolling_dividend"], 0.6)
        self.assertEqual(result["stocks"][0]["change_pct"], -0.012)
        self.assertEqual(result["stocks"][2]["code"], "00001")
        self.assertIsNone(result["stocks"][2]["quote_date"])
        self.assertIn("未记录真实交易日期", result["stocks"][2]["verification_note"])
        self.assertEqual(result["meta"]["fx_rate"], 0.8)
        self.assertEqual((self.root / "data" / "daily-dy.xlsx").read_bytes(), before)

    def test_mismatched_report_cannot_override_workbook_or_supply_history(self):
        self.report[0]["annual"]["total"] = 0.8
        self.report_path.write_text(json.dumps(self.report), encoding="utf-8")
        stock = extract(self.root)["stocks"][0]
        self.assertEqual(stock["annual_dividend"], 0.5)
        self.assertEqual(stock["history"], [])
        self.assertIsNone(stock["rolling_dividend"])
        self.assertIn("不一致", stock["verification_note"])

    def test_hk_dividend_follows_a_share_verification_and_current_amount(self):
        workbook_path = self.root / "data" / "daily-dy.xlsx"
        workbook = openpyxl.load_workbook(workbook_path)
        workbook["港股数据"].cell(2, 6).value = 0.8
        workbook["港股数据"].append(["待核实港股", "00002", "600002", 8, 8, 0.3, None, None, None, None, None, 0.8])
        workbook.save(workbook_path)
        workbook.close()
        stocks = extract(self.root)["stocks"]
        self.assertEqual(stocks[2]["annual_dividend"], 0.5)
        self.assertIsNone(stocks[3]["annual_dividend"])
        self.assertIn("对应 A 股年度股息待核实", stocks[3]["verification_note"])


if __name__ == "__main__":
    unittest.main()
