"""Live, read-only alternative-source probes restricted to workbook symbols."""
import concurrent.futures
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parents[1]


def worker(kind, symbol):
    from network_diagnostics import install
    install()
    import requests
    original = requests.Session.request
    def bounded(session, method, url, **kwargs):
        if kwargs.get('timeout') is None:
            kwargs['timeout'] = 12
        return original(session, method, url, **kwargs)
    requests.Session.request = bounded
    import akshare as ak
    started = time.perf_counter()
    try:
        if kind == 'cninfo':
            frame = ak.stock_dividend_cninfo(symbol=symbol)
        elif kind == 'xueqiu':
            if os.environ.get('XUEQIU_TOKEN'):
                from akshare.stock import cons
                cons.xq_a_token = os.environ['XUEQIU_TOKEN']
            frame = ak.stock_individual_spot_xq(symbol=symbol, timeout=12)
        elif kind == 'tencent':
            today = dt.date.today()
            frame = ak.stock_zh_a_hist_tx(symbol=symbol.lower(),
                    start_date=(today-dt.timedelta(days=21)).strftime('%Y%m%d'),
                    end_date=today.strftime('%Y%m%d'), adjust='', timeout=12)
        else:
            raise ValueError(kind)
        result = {'ok':True, 'rows':json.loads(frame.to_json(orient='records', date_format='iso'))}
    except Exception as exc:
        result = {'ok':False, 'error_type':type(exc).__name__, 'error':str(exc)}
    result.update(kind=kind, symbol=symbol, version=ak.__version__,
                  seconds=round(time.perf_counter()-started,3), fetched_at=dt.datetime.now().isoformat())
    print('PROBE_RESULT=' + json.dumps(result, ensure_ascii=False, allow_nan=False))


def main(tencent_only=False, xueqiu_only=False):
    import openpyxl
    path = ROOT / 'data/daily-dy.xlsx'
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    stamp = dt.datetime.now().strftime('%Y%m%d_%H%M%S')
    prefix = 'tencent_history_' if tencent_only else 'xueqiu_quotes_' if xueqiu_only else 'alternative_sources_'
    out = ROOT / 'pictures' / stamp[:8] / (prefix + stamp[9:])
    out.mkdir(parents=True, exist_ok=False)
    wb = openpyxl.load_workbook(path, data_only=True)
    stocks = []
    for row in wb['股票数据'].iter_rows(min_row=2, values_only=True):
        if not row[1] or not row[2]:
            continue
        code = str(row[2]).zfill(6)
        stocks.append(dict(market='A',name=row[1],code=code,
                           symbol=('SH' if code.startswith('6') else 'SZ')+code,
                           price=row[3],annual=row[4],annual_yield=row[6],
                           exdate=row[11],rolling_yield=row[12],year=row[13]))
    for row in wb['港股数据'].iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1] or not row[2]:
            continue
        code = str(row[1]).zfill(5)
        stocks.append(dict(market='HK',name=row[0],code=code,symbol=code,price=row[3]))
    wb.close()
    baseline = {'stocks':stocks,'workbook_sha256':before_hash,
                'dividend_details':json.loads((ROOT/'pictures'/stamp[:8]/'rolling_dividends.json').read_text(encoding='utf-8'))}
    (out/'baseline.json').write_text(json.dumps(baseline,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print('OUTPUT=' + str(out), flush=True)
    env = {**os.environ,'PYTHONUTF8':'1','DIVIDEND_DIAGNOSTICS_DIR':str(out)}
    if tencent_only:
        jobs = [('tencent',s,s['symbol'].lower()) for s in stocks if s['market']=='A']
    elif xueqiu_only:
        jobs = [('xueqiu',s,s['symbol']) for s in stocks]
    else:
        jobs = [('cninfo',s,s['code']) for s in stocks if s['market']=='A']
        jobs += [('xueqiu',s,s['symbol']) for s in stocks]
    def fetch(job):
        kind,stock,symbol = job
        try:
            p = subprocess.run([sys.executable,str(Path(__file__)), '--worker',kind,symbol],
                               capture_output=True,text=True,encoding='utf-8',errors='replace',
                               timeout=40,env=env)
            line = next(x for x in reversed(p.stdout.splitlines()) if x.startswith('PROBE_RESULT='))
            result = json.loads(line.split('=',1)[1])
        except Exception as exc:
            result = {'ok':False,'kind':kind,'symbol':symbol,'error_type':type(exc).__name__,'error':str(exc)}
        result.update(name=stock['name'],market=stock['market'],code=stock['code'])
        (out/f'{kind}_{stock["market"]}_{stock["code"]}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        print(f'{kind} {symbol} {stock["name"]}: ' + ('OK '+str(len(result['rows']))+' rows' if result['ok'] else result['error_type']+': '+result['error']),flush=True)
        return result
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        results = list(pool.map(fetch,jobs))
    unchanged = before_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    (out/'results.json').write_text(json.dumps({'results':results,'workbook_unchanged':unchanged},ensure_ascii=False,indent=2),encoding='utf-8')
    assert unchanged, 'Workbook changed during read-only test'
    print('DONE workbook unchanged; '+str(sum(r['ok'] for r in results))+'/'+str(len(results))+' calls succeeded',flush=True)


if __name__ == '__main__':
    if len(sys.argv)>1 and sys.argv[1]=='--worker':
        worker(sys.argv[2],sys.argv[3])
    else:
        main(tencent_only='--tencent' in sys.argv, xueqiu_only='--xueqiu' in sys.argv)
