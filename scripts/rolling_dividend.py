"""Refresh latest disclosed fiscal-year dividends and rolling annual yields."""
import calendar
import math
import re
import datetime as dt
import json
import os
from pathlib import Path
import subprocess
import copy
import zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
API = "https://datacenter-web.eastmoney.com/api/data/v1/get"
MISSING_INSTALLMENTS = {"000538": {"2025-06-30": {
    "REPORT_DATE": "2025-06-30", "ASSIGN_PROGRESS": "实施分配",
    "PRETAX_BONUS_RMB": 10.19, "EX_DIVIDEND_DATE": "2025-09-24"}}}


def supplement_missing(items, code):
    """Legacy correction scoped by security and period; never add it twice."""
    periods = {str(item.get('REPORT_DATE', ''))[:10] for item in items}
    return list(items) + [item for period, item in MISSING_INSTALLMENTS.get(code, {}).items()
                          if period not in periods]


def installment_dps(item):
    """Missing cash amounts are unavailable, not an implied zero dividend."""
    from update_dy import _parse_dps
    profile = str(item.get('IMPL_PLAN_PROFILE') or '')
    explicit_zero = '不分配' in str(item.get('ASSIGN_PROGRESS') or '')
    if item.get('PRETAX_BONUS_RMB') is None and not re.search(r'10派\d', profile) and not explicit_zero:
        raise ValueError('分红方案缺少每股金额，需人工核实')
    value = _parse_dps(item)
    if not math.isfinite(value) or value < 0:
        raise ValueError('分红金额无效，需人工核实')
    return value


def merge_authored_column(path, authored):
    """Keep original workbook parts; copy only affected Artifact-authored cells.

    Artifact import/export omits legacy Excel notes, so copying a narrow column
    avoids changing the original notes, metadata, styles or other worksheets.
    """
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    tag = lambda name: f"{{{ns}}}{name}"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(authored) as generated:
        native_styles=ET.fromstring(source.read('xl/styles.xml'))
        native_xfs=native_styles.find(tag('cellXfs'))
        format_styles={}
        strings = []
        if "xl/sharedStrings.xml" in generated.namelist():
            strings = ["".join(node.itertext()) for node in ET.fromstring(generated.read("xl/sharedStrings.xml"))]
        edits = {}
        has_history = 'xl/worksheets/sheet3.xml' in generated.namelist()
        for sheet_index, columns in [(1, {"E", "F", "G", "I", "J", "K", "M", "N"}),
                                     (2, {"F", "G", "H", "I", "J", "K"})]:
            if sheet_index == 1 and has_history:
                columns.update({'O','P','Q','R','S'})
            member = f"xl/worksheets/sheet{sheet_index}.xml"
            root = ET.fromstring(source.read(member))
            new_root = ET.fromstring(generated.read(member))
            cells = {c.attrib["r"]: c for c in new_root.iter(tag("c"))
                     if ''.join(filter(str.isalpha, c.attrib["r"])) in columns}
            for row in root.find(tag("sheetData")):
                originals = {c.attrib.get("r"): c for c in row}
                for column in sorted(columns):
                    address = f'{column}{row.attrib["r"]}'
                    if address not in cells:
                        continue
                    cell = copy.deepcopy(cells[address])
                    template = originals.get(address)
                    if template is None:
                        style_column='G' if column in ['K','M','O','P','R'] else 'I' if column in ['J','Q','S'] else 'F'
                        template = originals.get(f'{style_column}{row.attrib["r"]}')
                    cell.attrib.pop("s", None)
                    if template is not None and "s" in template.attrib:
                        cell.set("s", template.attrib["s"])
                    if sheet_index==1 and column in ['J','K','P','Q','R','S'] and int(row.attrib['r'])>1:
                        base=int(cell.get('s','0'))
                        number_format='10' if column in ['K','P','R'] else '2'
                        key=(base,number_format)
                        if key not in format_styles:
                            style=copy.deepcopy(native_xfs[base])
                            style.set('numFmtId',number_format)
                            style.set('applyNumberFormat','1')
                            index=next((str(i) for i,xf in enumerate(native_xfs) if ET.tostring(xf)==ET.tostring(style)),None)
                            if index is None:
                                index=str(len(native_xfs))
                                native_xfs.append(style)
                                native_xfs.set('count',str(len(native_xfs)))
                            format_styles[key]=index
                        cell.set('s',format_styles[key])
                    if cell.attrib.get("t") == "s":
                        value = strings[int(cell.find(tag("v")).text)]
                        cell.set("t", "inlineStr")
                        for child in list(cell):
                            cell.remove(child)
                        ET.SubElement(ET.SubElement(cell, tag("is")), tag("t")).text = value
                    for old in list(row):
                        if old.attrib.get("r") == address:
                            row.remove(old)
                    row.append(cell)
                row[:] = sorted(row, key=lambda c: (len(''.join(filter(str.isalpha, c.attrib.get('r', '')))), c.attrib.get('r', '')))
            if sheet_index == 1:
                cols = root.find(tag("cols"))
                if cols is None:
                    cols = ET.Element(tag("cols"))
                    root.insert(list(root).index(root.find(tag("sheetData"))), cols)
                widths = [(5, 23), (6, 18), (7, 18), (13, 26), (14, 18)]
                if has_history:
                    widths.extend([(15,28),(16,25),(17,23),(18,25),(19,23)])
                for index, width in widths:
                    for col in list(cols):
                        if col.attrib.get("min") == str(index) and col.attrib.get("max") == str(index):
                            cols.remove(col)
                    ET.SubElement(cols, tag("col"), {"min":str(index), "max":str(index), "width":str(width), "customWidth":"1"})
                last_row = max(int(row.attrib["r"]) for row in root.find(tag("sheetData")))
                dimension=root.find(tag('dimension'))
                old_end=dimension.get('ref','A1').split(':')[-1]
                old_column=re.match(r'[A-Z]+',old_end)[0]
                last_column=max(openpyxl.utils.column_index_from_string(old_column),19 if has_history else 14)
                dimension.set('ref',f'A1:{openpyxl.utils.get_column_letter(last_column)}{last_row}')
            edits[member] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        if has_history:
            # Add/replace only the supporting history tab, preserving native sheets/notes.
            rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
            content_ns = 'http://schemas.openxmlformats.org/package/2006/content-types'
            book = ET.fromstring(source.read('xl/workbook.xml'))
            rels = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
            content = ET.fromstring(source.read('[Content_Types].xml'))
            sheets = book.find(tag('sheets'))
            existing = next((s for s in sheets if s.get('name')=='股息率历史'),None)
            if existing is None:
                indices = [int(re.search(r'sheet(\d+)\.xml$',name)[1]) for name in source.namelist()
                           if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name)]
                member = f'xl/worksheets/sheet{max(indices,default=0)+1}.xml'
                used_ids = {r.get('Id') for r in rels}
                i=1
                while f'rId{i}' in used_ids:
                    i+=1
                rid=f'rId{i}'
                ET.SubElement(sheets,tag('sheet'),{'name':'股息率历史','sheetId':str(max(int(s.get('sheetId')) for s in sheets)+1),
                    '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id':rid})
                ET.SubElement(rels,f'{{{rel_ns}}}Relationship',{'Id':rid,'Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet','Target':member[3:]})
                ET.SubElement(content,f'{{{content_ns}}}Override',{'PartName':'/'+member,'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
            else:
                rid=existing.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                target=next(r.get('Target') for r in rels if r.get('Id')==rid)
                member=target.lstrip('/') if target.startswith('/') else 'xl/'+target
            history=ET.fromstring(generated.read('xl/worksheets/sheet3.xml'))
            styles=native_styles
            cell_xfs=styles.find(tag('cellXfs'))
            main=ET.fromstring(source.read('xl/worksheets/sheet1.xml'))
            source_styles={c.get('r'):c.get('s','0') for c in main.iter(tag('c'))}
            percent_style=source_styles.get('G2','0')
            header_style=source_styles.get('G1','0')
            formats=styles.find(tag('numFmts'))
            if formats is None:
                formats=ET.Element(tag('numFmts'),{'count':'0'})
                styles.insert(0,formats)
            date_format=next((n.get('numFmtId') for n in formats if n.get('formatCode')=='yyyy-mm-dd'),None)
            if date_format is None:
                date_format=str(max([163]+[int(n.get('numFmtId')) for n in formats])+1)
                ET.SubElement(formats,tag('numFmt'),{'numFmtId':date_format,'formatCode':'yyyy-mm-dd'})
                formats.set('count',str(len(formats)))
            date_xf=copy.deepcopy(cell_xfs[int(percent_style)])
            date_xf.set('numFmtId',date_format)
            date_xf.set('applyNumberFormat','1')
            date_style=next((str(i) for i,xf in enumerate(cell_xfs) if ET.tostring(xf)==ET.tostring(date_xf)),None)
            if date_style is None:
                date_style=str(len(cell_xfs))
                cell_xfs.append(date_xf)
                cell_xfs.set('count',str(len(cell_xfs)))
            text_style=next((str(i) for i,xf in enumerate(cell_xfs) if xf.get('numFmtId')=='49'),None)
            if text_style is None:
                text_xf=copy.deepcopy(cell_xfs[0])
                text_xf.set('numFmtId','49')
                text_xf.set('applyNumberFormat','1')
                text_style=str(len(cell_xfs))
                cell_xfs.append(text_xf)
                cell_xfs.set('count',str(len(cell_xfs)))
            for cell in history.iter(tag('c')):
                address=cell.get('r')
                column=re.match(r'[A-Z]+',address)[0]
                number=0
                for char in column:
                    number=number*26+ord(char)-64
                row=int(re.search(r'\d+',address)[0])
                style=header_style if row==5 else text_style if row>=6 and number==1 else date_style if row>=6 and number in [3,4] else percent_style if row>=6 and number in [5,8,9] else '0'
                cell.set('s',style)
                if cell.get('t')=='s':
                    value=strings[int(cell.find(tag('v')).text)]
                    cell.set('t','inlineStr')
                    for child in list(cell):cell.remove(child)
                    ET.SubElement(ET.SubElement(cell,tag('is')),tag('t')).text=value
            edits[member]=ET.tostring(history,encoding='utf-8',xml_declaration=True)
            edits['xl/workbook.xml']=ET.tostring(book,encoding='utf-8',xml_declaration=True)
            ET.register_namespace('',rel_ns)
            edits['xl/_rels/workbook.xml.rels']=ET.tostring(rels,encoding='utf-8',xml_declaration=True)
            ET.register_namespace('',content_ns)
            edits['[Content_Types].xml']=ET.tostring(content,encoding='utf-8',xml_declaration=True)
            ET.register_namespace('',ns)
        edits['xl/styles.xml']=ET.tostring(native_styles,encoding='utf-8',xml_declaration=True)
        temp = path.with_suffix(".tmp.xlsx")
        with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as out:
            for info in source.infolist():
                out.writestr(info, edits.get(info.filename, source.read(info.filename)))
            for name,data in edits.items():
                if name not in source.namelist():
                    out.writestr(name,data)
    temp.replace(path)


def eligible_periods(items, today=None):
    today = today or dt.date.today()
    periods = {}
    for item in items:
        progress = str(item.get("ASSIGN_PROGRESS") or "")
        implemented = any(k in progress for k in ("实施分配", "特别分配", "不分配"))
        if not implemented and not any(k in progress for k in (
                "董事会决议通过", "股东大会决议通过", "股东大会通过")):
            continue
        try:
            period = dt.date.fromisoformat(str(item["REPORT_DATE"])[:10])
        except (ValueError, KeyError):
            continue
        if period > today:
            continue
        # One current plan per reporting period; implemented supersedes proposed.
        existing = periods.get(period)
        if existing is None or (implemented and "实施分配" not in str(existing.get("ASSIGN_PROGRESS"))):
            periods[period] = item
    if not periods:
        raise ValueError("无已公告或已实施的分红记录")
    return periods


def select_annual(items, today=None):
    """Latest disclosed full fiscal year; never infer a year from payment date."""
    periods = eligible_periods(items, today)
    annual_dates = [p for p in periods if (p.month, p.day) == (12, 31)]
    if not annual_dates:
        raise ValueError("尚无已披露的全年分红方案，不能将中报金额当作全年")
    year = max(annual_dates).year
    if (today or dt.date.today()).year - year > 2:
        raise ValueError("全年分红报告期过旧，需人工核实")
    parts = []
    for period, item in sorted(periods.items()):
        if period.year != year:
            continue
        implemented = any(k in item["ASSIGN_PROGRESS"] for k in ["实施分配", "特别分配", "不分配"])
        parts.append({"period": period.isoformat(), "dps": installment_dps(item),
                      "status": item["ASSIGN_PROGRESS"], "implemented": implemented,
                      "date": str(item.get("EX_DIVIDEND_DATE") or period)[:10]})
    return {"year": year, "total": round(sum(p["dps"] for p in parts), 8),
            "confirmed": int(all(p["implemented"] for p in parts)), "parts": parts}


def select_rolling(items, today=None):
    """Use announced/implemented installments in (latest period - 1 year, latest]."""
    today = today or dt.date.today()
    periods = eligible_periods(items, today)
    latest = max(periods)
    cutoff = latest.replace(year=latest.year - 1,
                            day=min(latest.day, calendar.monthrange(latest.year - 1, latest.month)[1]))
    if (today - latest).days > 550:
        raise ValueError("最新分红报告期过旧，需人工核实")
    parts = [{"period": period.isoformat(), "dps": installment_dps(item),
              "status": item["ASSIGN_PROGRESS"]}
             for period, item in sorted(periods.items()) if cutoff < period <= latest]
    return {"total": round(sum(part["dps"] for part in parts), 8),
            "period_end": latest.isoformat(), "parts": parts}


def fetch_one(stock):
    try:
        from data_sources import dividend_records
        def fallback():
            params = dict(reportName="RPT_SHAREBONUS_DET", columns="ALL",
                          filter=f'(SECURITY_CODE="{stock["code"]}")', pageNumber=1,
                          pageSize=50, sortColumns="REPORT_DATE", sortTypes=-1)
            response = requests.get(API, params=params, timeout=25)
            response.raise_for_status()
            items = (response.json().get("result") or {}).get("data")
            if items is None:
                raise ValueError("接口未返回数据")
            return items, response.url
        items, source = dividend_records(stock['code'], fallback)
        items = supplement_missing(items, stock['code'])
        result = {**stock, **select_rolling(items), "source": source}
        try:
            result["annual"] = select_annual(items)
        except ValueError as exc:
            result["annual_error"] = str(exc)
        return result
    except Exception as exc:
        return {**stock, "error": str(exc)}


def authoring_input(path,out):
    """Project source metadata without importing the history table we rebuild."""
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rid_name='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    with zipfile.ZipFile(path) as source:
        book=ET.fromstring(source.read('xl/workbook.xml'))
        sheet=next((s for s in book.find('s:sheets',ns) if s.get('name')=='股息率历史'),None)
        if sheet is None:
            return path
        rels=ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        target=next(r.get('Target') for r in rels if r.get('Id')==sheet.get(rid_name))
        member=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        projected=out/'rolling_input.xlsx'
        with zipfile.ZipFile(projected,'w',zipfile.ZIP_DEFLATED) as dest:
            for info in source.infolist():
                data=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>' if info.filename==member else source.read(info.filename)
                dest.writestr(info,data)
    return projected


def refresh(records=None, percentile_as_of=None):
    path = ROOT / "data/daily-dy.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["股票数据"]
    stocks = [{"row": r, "name": ws.cell(r, 2).value,
               "code": str(ws.cell(r, 3).value).zfill(6)}
              for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value and ws.cell(r, 3).value]
    year_start_header=ws.cell(1,10).value
    year_start_values={s['code']:ws.cell(s['row'],10).value for s in stocks}
    wb.close()
    if records is None:
        with ThreadPoolExecutor(max_workers=6) as pool:
            records = list(pool.map(fetch_one, stocks))
    from year_start import prepare as prepare_year_start
    records=prepare_year_start(records,year_start_values,year_start_header)
    if percentile_as_of is not None:
        from history_percentile import prepare
        records = prepare(records,percentile_as_of)
    out = ROOT / "pictures" / dt.date.today().strftime("%Y%m%d")
    out.mkdir(parents=True, exist_ok=True)
    data_path = out / "rolling_dividends.json"
    data_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    for record in records:
        print(f'  {record["name"]}: {record.get("total", record.get("error"))}', flush=True)
    runtime = Path(os.environ.get("USERPROFILE", str(Path.home()))) / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
    input_path=authoring_input(path,out)
    subprocess.run([str(runtime), str(ROOT / "scripts/write_rolling_column.mjs"),
                    str(input_path), str(data_path)], check=True, cwd=ROOT)
    merge_authored_column(path, out / "rolling_authored.xlsx")
    failures = [r for r in records if "error" in r]
    print(f'滚动股息率：{len(records)-len(failures)}/{len(records)} 支成功；失败项显示待核实。', flush=True)
    return records


if __name__ == "__main__":
    refresh()
