from pathlib import Path
import sys
import unittest
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import openpyxl
import update_dy


class MissingTargetYieldTest(unittest.TestCase):
    def test_missing_target_does_not_use_text_or_numeric_target_price_as_gap(self):
        for target in [None,'待核实']:
            for old_price in ['待核实',100]:
                with self.subTest(target=target,old_price=old_price):
                    wb=openpyxl.Workbook();ws=wb.active;ws.title='股票数据'
                    ws.cell(2,2,'云铝股份');ws.cell(2,3,'000807');ws.cell(2,4,10)
                    ws.cell(2,5,.699);ws.cell(2,8,target);ws.cell(2,9,old_price);ws.cell(2,14,2025)
                    with patch.object(update_dy,'CACHE',{}),patch.object(update_dy,'_fetch_ashare_div',return_value=(None,None)), \
                         patch('data_sources.cached_dividends',return_value=([],'this run')):
                        update_dy.update_ashare_sheet(wb,wb,{'000807':11})
                    self.assertEqual(ws.cell(2,4).value,11)
                    self.assertEqual(ws.cell(2,8).value,target)
                    self.assertEqual(ws.cell(2,2).fill,update_dy.FILL_CLEAR)

    def test_invalid_gap_has_neutral_fill(self):
        for value in [None,'待核实','',float('nan'),float('inf'),True]:
            with self.subTest(value=value):self.assertEqual(update_dy.gh_fill(value),update_dy.FILL_CLEAR)

    def test_numeric_gap_keeps_existing_color_bands(self):
        for value,expected in [(0,update_dy.FILL_GREEN),(-.001,update_dy.FILL_BLUE),
                               (-.0028,update_dy.FILL_ORANGE),(-.01,update_dy.FILL_CLEAR)]:
            with self.subTest(value=value):self.assertEqual(update_dy.gh_fill(value),expected)
